"""Unit tests for ingestion handlers, exercising the example fixtures in
``data/upload_examples/``. Each handler is tested independently — no DB,
no ES, no LLM."""
import os
from pathlib import Path

import pytest


@pytest.fixture(scope="module")
def fixtures_dir() -> Path:
    p = Path(__file__).resolve().parents[2] / "data" / "upload_examples"
    assert p.is_dir(), f"fixtures dir missing: {p}"
    return p


# ── CsvHandler ────────────────────────────────────────────────────────────────

class TestCsvHandler:

    @pytest.fixture
    def handler(self):
        from src.ingestion.handlers.csv_handler import CsvHandler
        return CsvHandler()

    def test_sniff_recognizes_csv(self, handler, fixtures_dir):
        head = (fixtures_dir / "csv" / "leak_users_2024.csv").read_bytes()[:4096]
        assert handler.sniff(head) > 0.5

    def test_sniff_rejects_random_text(self, handler):
        assert handler.sniff(b"this is just prose with no commas\nand no tabs\n") == 0.0

    def test_sniff_rejects_binary(self, handler):
        assert handler.sniff(b"\x00\x01\x02hello world") == 0.0

    def test_extract_yields_one_record_per_row(self, handler, fixtures_dir):
        path = fixtures_dir / "csv" / "leak_users_2024.csv"
        proposal = handler.propose_mapping(path, {})
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        # leak_users_2024.csv has 10 data rows
        assert len(records) == 10
        # Every record carries the email and ip in raw
        first = records[0]
        assert "email" in first.raw
        assert "ip_address" in first.raw
        assert "email:" in first.text
        assert "ip_address:" in first.text

    def test_extract_wide_csv_preserves_and_searches_all_columns(self, handler, tmp_path):
        headers = [f"col_{i}" for i in range(100)]
        values = [f"value_{i}" for i in range(100)]
        path = tmp_path / "wide.csv"
        path.write_text(",".join(headers) + "\n" + ",".join(values) + "\n", encoding="utf-8")

        records = list(handler.extract(path, {"text": "col_0"}, {"delimiter": ",", "encoding": "utf-8"}))

        assert len(records) == 1
        assert len(records[0].raw) == 100
        assert records[0].raw["col_99"] == "value_99"
        assert "col_99: value_99" in records[0].text

    def test_propose_mapping_finds_text_column(self, handler, fixtures_dir):
        path = fixtures_dir / "csv" / "support_tickets.csv"
        proposal = handler.propose_mapping(path, {})
        # `body` is the longest column → should be picked for text
        assert proposal.suggested_mapping["text"] == "body"
        # `subject` should be picked as title
        assert proposal.suggested_mapping["title"] == "subject"
        # `created_at` should be detected
        assert proposal.suggested_mapping["created_at"] == "created_at"

    def test_propose_mapping_handles_semicolon_delimiter(self, handler, fixtures_dir):
        # intel_reports.csv uses `;` as delimiter
        path = fixtures_dir / "csv" / "intel_reports.csv"
        proposal = handler.propose_mapping(path, {})
        assert proposal.options["delimiter"] == ";"
        assert "summary" in proposal.columns
        # `summary` is the longest column → picked as text
        assert proposal.suggested_mapping["text"] == "summary"

    def test_fingerprint_stable_across_runs(self, handler, fixtures_dir):
        path = fixtures_dir / "csv" / "leak_users_2024.csv"
        fp1 = handler.fingerprint(path, {})
        fp2 = handler.fingerprint(path, {})
        assert fp1 == fp2 and fp1 is not None

    def test_fingerprint_matches_for_same_structure(self, handler, fixtures_dir):
        """leak_users_2024 and leak_users_2025 share headers — same fingerprint."""
        fp_2024 = handler.fingerprint(fixtures_dir / "csv" / "leak_users_2024.csv", {})
        fp_2025 = handler.fingerprint(fixtures_dir / "csv" / "leak_users_2025.csv", {})
        assert fp_2024 == fp_2025

    def test_fingerprint_differs_for_different_structure(self, handler, fixtures_dir):
        fp_users   = handler.fingerprint(fixtures_dir / "csv" / "leak_users_2024.csv", {})
        fp_tickets = handler.fingerprint(fixtures_dir / "csv" / "support_tickets.csv", {})
        assert fp_users != fp_tickets


# ── JsonHandler ───────────────────────────────────────────────────────────────

class TestJsonHandler:

    @pytest.fixture
    def handler(self):
        from src.ingestion.handlers.json_handler import JsonHandler
        return JsonHandler()

    def test_sniff_recognizes_jsonl(self, handler, fixtures_dir):
        head = (fixtures_dir / "jsonl" / "tweets.jsonl").read_bytes()[:4096]
        assert handler.sniff(head) > 0.5

    def test_extract_jsonl(self, handler, fixtures_dir):
        path = fixtures_dir / "jsonl" / "tweets.jsonl"
        proposal = handler.propose_mapping(path, {})
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        # tweets.jsonl has 8 entries
        assert len(records) == 8
        # text column should be `text`
        assert proposal.suggested_mapping["text"] == "text"
        # author should be detected
        assert proposal.suggested_mapping["author"] == "author"

    def test_extract_json_array(self, handler, fixtures_dir):
        path = fixtures_dir / "misc" / "threat_actors.json"
        proposal = handler.propose_mapping(path, {})
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        # threat_actors.json has 4 actors
        assert len(records) == 4

    def test_fingerprint_jsonl_uses_keys(self, handler, fixtures_dir):
        fp_a = handler.fingerprint(fixtures_dir / "jsonl" / "tweets.jsonl", {})
        fp_b = handler.fingerprint(fixtures_dir / "jsonl" / "osint_findings.jsonl", {})
        # Different schemas, different fingerprints
        assert fp_a is not None and fp_b is not None
        assert fp_a != fp_b


# ── TextHandler ───────────────────────────────────────────────────────────────

class TestTextHandler:

    @pytest.fixture
    def handler(self):
        from src.ingestion.handlers.text_handler import TextHandler
        return TextHandler()

    def test_telegram_chat_pattern_detected(self, handler, fixtures_dir):
        path = fixtures_dir / "txt" / "chat_log_telegram.txt"
        proposal = handler.propose_mapping(path, {})
        assert proposal.options["mode"] == "regex"
        assert proposal.options.get("regex_name") == "telegram_chat"
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        # Each line in the chat is one record
        assert len(records) == 12
        # First record's author starts with @
        assert records[0].author.startswith("@")
        assert records[0].created_at is not None

    def test_apache_log_pattern_detected(self, handler, fixtures_dir):
        path = fixtures_dir / "txt" / "server_access.log"
        proposal = handler.propose_mapping(path, {})
        assert proposal.options["mode"] == "regex"
        assert proposal.options.get("regex_name") == "apache_access"
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        assert len(records) >= 10
        assert all(r.created_at for r in records)

    def test_paragraph_mode_for_prose(self, handler, fixtures_dir):
        path = fixtures_dir / "txt" / "incident_report.txt"
        proposal = handler.propose_mapping(path, {})
        # Prose with section headings — should fall back to paragraph mode
        assert proposal.options["mode"] == "paragraph"
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        assert len(records) >= 4   # several paragraphs

    def test_freeform_notes_paragraph(self, handler, fixtures_dir):
        path = fixtures_dir / "txt" / "notes_freeform.txt"
        proposal = handler.propose_mapping(path, {})
        assert proposal.options["mode"] in ("paragraph", "line")
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        assert len(records) > 0


# ── XlsxHandler smoke (no fixture file — write one in test) ──────────────────

class TestXlsxHandler:

    @pytest.fixture
    def handler(self):
        from src.ingestion.handlers.xlsx_handler import XlsxHandler
        return XlsxHandler()

    def test_xlsx_extract(self, handler, tmp_path):
        # Build a tiny xlsx in memory and ingest it.
        from openpyxl import Workbook
        path = tmp_path / "sample.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "title", "body", "created_at"])
        ws.append([1, "First entry", "This is the body of the first entry — long text.", "2024-01-01"])
        ws.append([2, "Second entry", "The body of the second entry, also reasonably long.", "2024-01-02"])
        wb.save(path)
        wb.close()

        proposal = handler.propose_mapping(path, {})
        # `body` is the longest column → picked for text
        assert proposal.suggested_mapping["text"] == "body"
        records = list(handler.extract(path, proposal.suggested_mapping, proposal.options))
        assert len(records) == 2
        assert records[0].title == "First entry"

    def test_xlsx_multi_sheet_ingests_every_sheet(self, handler, tmp_path):
        """Two-sheet workbook → every row from every sheet must come out
        (no silent data loss)."""
        from openpyxl import Workbook
        path = tmp_path / "multi.xlsx"
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Alpha"
        ws1.append(["id", "body"])
        ws1.append([1, "alpha row 1"])
        ws1.append([2, "alpha row 2"])
        ws2 = wb.create_sheet("Beta")
        ws2.append(["id", "body"])
        ws2.append([10, "beta row 1"])
        ws2.append([20, "beta row 2"])
        ws2.append([30, "beta row 3"])
        wb.save(path)
        wb.close()

        records = list(handler.extract(
            path, {"text": "body"}, {"has_header": True},
        ))
        # 2 + 3 = 5 records across both sheets
        assert len(records) == 5
        sheets = {r.raw.get("__sheet") for r in records}
        assert sheets == {"Alpha", "Beta"}
        # record_index is contiguous (no collisions between sheets)
        indices = [r.record_index for r in records]
        assert indices == [0, 1, 2, 3, 4]
